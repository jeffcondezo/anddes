from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any, BinaryIO

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from maestro.models import Criterio, Documento, DocumentoCriterio, RolDocumentoCriterio

SHEET_NAME = "Documento_x_Criterio"
INSTRUCTIONS_SHEET = "Instrucciones"

HEADERS = (
    "ID_Documento",
    "Documento",
    "ID_Criterio",
    "Rol",
    "ID_Documento_Secundario",
    "Principio",
    "Criterio",
)

# Columnas 0-based que el importador sí lee.
REQUIRED_COLS = {0, 2}  # ID_Documento, ID_Criterio
USED_OPTIONAL_COLS = {3, 4}  # Rol, ID_Documento_Secundario
# 1, 5, 6 = solo referencia (Documento nombre, Principio, Criterio texto)

FILL_REQUIRED = PatternFill("solid", fgColor="C6EFCE")  # verde
FILL_USED = PatternFill("solid", fgColor="FFEB9C")  # ámbar
FILL_REF = PatternFill("solid", fgColor="D9D9D9")  # gris
FONT_HEADER = Font(bold=True, color="000000")
FONT_TITLE = Font(bold=True, size=14)
FONT_SECTION = Font(bold=True, size=11)


def _style_data_header(ws) -> None:
    for idx, title in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=idx, value=title)
        cell.font = FONT_HEADER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        col0 = idx - 1
        if col0 in REQUIRED_COLS:
            cell.fill = FILL_REQUIRED
            cell.comment = Comment(
                "OBLIGATORIO al importar. Debe existir en el catálogo del sistema.",
                "GISTM",
                width=280,
                height=60,
            )
        elif col0 in USED_OPTIONAL_COLS:
            cell.fill = FILL_USED
            if col0 == 3:
                cell.comment = Comment(
                    "Se usa al importar. Valores: PRINCIPAL o COMPLEMENTARIO. "
                    "Si está vacío, se asume PRINCIPAL.",
                    "GISTM",
                    width=300,
                    height=70,
                )
            else:
                cell.comment = Comment(
                    "Se usa al importar (opcional). Código del documento secundario. "
                    "Si no existe en el catálogo, la fila se guarda sin secundario.",
                    "GISTM",
                    width=300,
                    height=80,
                )
        else:
            cell.fill = FILL_REF
            cell.comment = Comment(
                "Solo referencia / ayuda visual. NO se usa al importar.",
                "GISTM",
                width=260,
                height=50,
            )
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"


def _build_instructions_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet(INSTRUCTIONS_SHEET, 0)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 72

    rows: list[tuple[str, str]] = [
        ("Relación estándar · Instrucciones", ""),
        ("", ""),
        (
            "Para qué sirve este archivo",
            "Es un backup editable de las relaciones Documento ↔ Criterio del estándar. "
            "Puedes descargarlo, modificarlo en Excel y volver a cargarlo en el sistema.",
        ),
        (
            "Flujo recomendado",
            "1) Descargar este Excel.  2) Editar las filas que necesites.  "
            "3) En el sistema: Borrar relaciones actuales.  4) Importar este archivo.",
        ),
        ("", ""),
        ("Leyenda de colores (hoja de datos)", ""),
        ("Verde", "Obligatorio al importar."),
        ("Ámbar / amarillo", "Se usa al importar, pero puede ir vacío."),
        ("Gris", "Solo referencia; el importador lo ignora."),
        ("", ""),
        ("Campos de la hoja Documento_x_Criterio", ""),
        ("ID_Documento (col. A) · verde", "Obligatorio. Código del documento principal (ej. D001). Debe existir en el catálogo."),
        ("Documento (col. B) · gris", "Nombre del documento. Solo ayuda visual; no se importa."),
        ("ID_Criterio (col. C) · verde", "Obligatorio. Código del criterio (ej. C001). Debe existir en el catálogo."),
        (
            "Rol (col. D) · ámbar",
            "PRINCIPAL o COMPLEMENTARIO. Si falta, se asume PRINCIPAL.",
        ),
        (
            "ID_Documento_Secundario (col. E) · ámbar",
            "Opcional. Segunda fuente de evidencia (típico en Complementario). "
            "Si el código no existe, la relación se guarda sin secundario.",
        ),
        ("Principio (col. F) · gris", "Solo referencia; no se importa."),
        ("Criterio (col. G) · gris", "Texto/descripción de ayuda; no se importa."),
        ("", ""),
        ("Si un ID no existe", ""),
        (
            "Documento o criterio inexistente",
            "Esa fila se omite. El resto se importa. Al final verás un mensaje con "
            "creados / actualizados / omitidos y la lista de códigos no encontrados.",
        ),
        (
            "Secundario inexistente",
            "La fila sí se importa, pero sin documento secundario.",
        ),
        ("", ""),
        ("Notas", ""),
        (
            "Hoja requerida",
            f"El archivo debe contener la hoja «{SHEET_NAME}». Esta hoja de instrucciones puede quedarse.",
        ),
        (
            "Un criterio = una fila",
            "Cada criterio solo puede aparecer una vez. Si se repite, prevalece la última fila leída al importar.",
        ),
        (
            "Nombres del catálogo",
            "Los nombres de documento/criterio los resuelve el sistema por el código; "
            "no hace falta que el Excel traiga el nombre exacto.",
        ),
    ]

    for r_idx, (left, right) in enumerate(rows, start=1):
        c1 = ws.cell(row=r_idx, column=1, value=left)
        c2 = ws.cell(row=r_idx, column=2, value=right)
        c2.alignment = Alignment(wrap_text=True, vertical="top")
        if r_idx == 1:
            c1.font = FONT_TITLE
        elif left and not right and left not in ("",):
            c1.font = FONT_SECTION
        elif left == "Verde":
            c1.fill = FILL_REQUIRED
            c1.font = FONT_HEADER
        elif left.startswith("Ámbar"):
            c1.fill = FILL_USED
            c1.font = FONT_HEADER
        elif left == "Gris":
            c1.fill = FILL_REF
            c1.font = FONT_HEADER
        elif "· verde" in left:
            c1.fill = FILL_REQUIRED
        elif "· ámbar" in left:
            c1.fill = FILL_USED
        elif "· gris" in left:
            c1.fill = FILL_REF

    ws.row_dimensions[3].height = 45
    ws.row_dimensions[4].height = 45
    for r in range(12, 19):
        ws.row_dimensions[r].height = 36
    ws.row_dimensions[21].height = 45
    ws.row_dimensions[22].height = 30


def build_plantilla_workbook() -> Workbook:
    """
    Exporta las relaciones estándar actuales en el mismo formato que acepta
    la importación (backup editable), con hoja de instrucciones.
    """
    wb = Workbook()
    # Reemplazar la hoja por defecto: primero instrucciones, luego datos.
    default = wb.active
    wb.remove(default)
    _build_instructions_sheet(wb)

    ws = wb.create_sheet(SHEET_NAME)
    _style_data_header(ws)

    relaciones = (
        DocumentoCriterio.objects.select_related(
            "documento",
            "documento_secundario",
            "criterio",
            "criterio__principio",
        )
        .order_by("criterio__codigo", "documento__codigo")
    )
    for rel in relaciones:
        principio = ""
        if rel.criterio.principio_id:
            principio = rel.criterio.principio.codigo
        ws.append(
            [
                rel.documento.codigo,
                rel.documento.nombre,
                rel.criterio.codigo,
                rel.rol,
                rel.documento_secundario.codigo if rel.documento_secundario_id else "",
                principio,
                (rel.criterio.descripcion or "")[:200],
            ]
        )

    for idx, width in enumerate((14, 40, 14, 16, 22, 14, 48), start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    return wb


def plantilla_bytes() -> bytes:
    wb = build_plantilla_workbook()
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@transaction.atomic
def importar_documento_criterio_excel(source: str | Path | BinaryIO) -> dict[str, Any]:
    """
    Importa relaciones Documento↔Criterio desde Documento_x_Criterio.
    Columnas: A ID_Documento, C ID_Criterio, D Rol, E ID_Documento_Secundario.
    """
    wb = load_workbook(source, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"La hoja '{SHEET_NAME}' no existe en el Excel.")
    ws = wb[SHEET_NAME]

    docs = {d.codigo: d for d in Documento.objects.all()}
    crits = {c.codigo: c for c in Criterio.objects.all()}

    created = 0
    existing = 0
    skipped = 0
    missing_docs: set[str] = set()
    missing_crits: set[str] = set()

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        id_doc = str(row[0]).strip() if row[0] is not None else ""
        id_crit = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
        rol_raw = str(row[3]).strip().upper() if len(row) > 3 and row[3] else "PRINCIPAL"
        id_sec = ""
        if len(row) > 4 and row[4] is not None:
            id_sec = str(row[4]).strip()
            if id_sec in ("", "-", "None", "none"):
                id_sec = ""
        if not id_doc or not id_crit:
            continue
        documento = docs.get(id_doc)
        criterio = crits.get(id_crit)
        if documento is None:
            missing_docs.add(id_doc)
            skipped += 1
            continue
        if criterio is None:
            missing_crits.add(id_crit)
            skipped += 1
            continue
        doc_sec = None
        if id_sec:
            doc_sec = docs.get(id_sec)
            if doc_sec is None:
                missing_docs.add(id_sec)
            elif doc_sec.pk == documento.pk:
                doc_sec = None

        if rol_raw.startswith("COMP"):
            rol = RolDocumentoCriterio.COMPLEMENTARIO
        else:
            rol = RolDocumentoCriterio.PRINCIPAL

        _, was_created = DocumentoCriterio.objects.update_or_create(
            criterio=criterio,
            defaults={
                "documento": documento,
                "documento_secundario": doc_sec,
                "rol": rol,
            },
        )
        if was_created:
            created += 1
        else:
            existing += 1

    return {
        "created": created,
        "existing": existing,
        "skipped": skipped,
        "missing_docs": sorted(missing_docs),
        "missing_crits": sorted(missing_crits),
    }


def format_import_message(stats: dict[str, Any]) -> str:
    parts = [
        f"creados={stats['created']}",
        f"actualizados={stats['existing']}",
        f"omitidos={stats['skipped']}",
    ]
    msg = "Importación lista: " + ", ".join(parts)
    if stats["missing_docs"]:
        docs = ", ".join(stats["missing_docs"][:15])
        if len(stats["missing_docs"]) > 15:
            docs += "..."
        msg += f". Documentos no encontrados: {docs}"
    if stats["missing_crits"]:
        crits = ", ".join(stats["missing_crits"][:15])
        if len(stats["missing_crits"]) > 15:
            crits += "..."
        msg += f". Criterios no encontrados: {crits}"
    return msg
