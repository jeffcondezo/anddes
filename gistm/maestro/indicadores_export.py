from __future__ import annotations

from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .indicadores import DocumentoIndicador


def _fecha(val) -> str:
    if not val:
        return "—"
    try:
        return val.strftime("%d/%m/%Y")
    except Exception:
        return str(val)


def _filas(docs: list[DocumentoIndicador]) -> list[list[str]]:
    rows = [
        [
            "Código",
            "Nombre",
            "Origen",
            "Estado",
            "Áreas",
            "Principios",
            "Revisado por",
            "Fecha límite",
            "Fecha carga",
            "Fecha revisión",
        ]
    ]
    for d in docs:
        rows.append(
            [
                d.codigo,
                d.nombre,
                d.origen_label,
                d.estado_label,
                ", ".join(d.areas) if d.areas else "—",
                ", ".join(d.principios) if d.principios else "—",
                d.revisado_por or "—",
                _fecha(d.fecha_limite),
                _fecha(d.fecha_carga),
                _fecha(d.fecha_revision),
            ]
        )
    return rows


def exportar_documentos_excel(
    docs: list[DocumentoIndicador],
    *,
    empresa_nombre: str,
    revision_etiqueta: str,
) -> HttpResponse:
    wb = Workbook()
    ws = wb.active
    ws.title = "Documentos"
    ws.append([f"Documentos · {empresa_nombre} · {revision_etiqueta}"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    for row in _filas(docs):
        ws.append(row)
    for cell in ws[3]:
        cell.font = Font(bold=True)
    from openpyxl.utils import get_column_letter

    for idx in range(1, 11):
        letter = get_column_letter(idx)
        max_len = 12
        for row in ws.iter_rows(min_row=3, min_col=idx, max_col=idx):
            for cell in row:
                if cell.value:
                    max_len = max(max_len, min(len(str(cell.value)), 48))
        ws.column_dimensions[letter].width = max_len + 2
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = (
        f'attachment; filename="documentos_{empresa_nombre[:40]}.xlsx"'
    )
    return resp


def exportar_documentos_pdf(
    docs: list[DocumentoIndicador],
    *,
    empresa_nombre: str,
    revision_etiqueta: str,
) -> HttpResponse:
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "TituloDocs",
        parent=styles["Heading2"],
        fontSize=12,
        spaceAfter=8,
    )
    cell = ParagraphStyle(
        "CeldaDocs",
        parent=styles["Normal"],
        fontSize=7,
        leading=9,
    )
    story = [
        Paragraph(
            f"Documentos · {empresa_nombre} · {revision_etiqueta}",
            title,
        ),
        Spacer(1, 4),
    ]
    data = []
    for i, row in enumerate(_filas(docs)):
        if i == 0:
            data.append(row)
        else:
            data.append([Paragraph(str(c), cell) for c in row])
    table = Table(
        data,
        colWidths=[
            22 * mm,
            55 * mm,
            22 * mm,
            28 * mm,
            35 * mm,
            28 * mm,
            32 * mm,
            22 * mm,
            22 * mm,
            24 * mm,
        ],
        repeatRows=1,
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0069aa")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 8),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cfd8dc")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f7fafc")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    story.append(table)
    if not docs:
        story.append(Spacer(1, 8))
        story.append(Paragraph("Sin documentos para exportar.", styles["Normal"]))
    doc.build(story)
    buf.seek(0)
    resp = HttpResponse(buf.getvalue(), content_type="application/pdf")
    resp["Content-Disposition"] = (
        f'attachment; filename="documentos_{empresa_nombre[:40]}.pdf"'
    )
    return resp
