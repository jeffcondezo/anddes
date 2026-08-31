from __future__ import annotations

from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from maestro.models import Criterio, Documento, DocumentoCriterio


class Command(BaseCommand):
    help = (
        "Importa la relación estándar Documento↔Criterio desde la hoja "
        "Documento_x_Criterio del Excel (ID_Documento, ID_Criterio)."
    )

    def add_arguments(self, parser):
        parser.add_argument(
            "--path",
            default="",
            help="Ruta al excel.xlsx (por defecto gistm/excel.xlsx).",
        )

    @transaction.atomic
    def handle(self, *args, **options):
        path_opt = (options.get("path") or "").strip()
        if path_opt:
            excel_path = Path(path_opt)
        else:
            # manage.py vive en gistm/; excel en gistm/excel.xlsx
            excel_path = Path(__file__).resolve().parents[3] / "excel.xlsx"

        if not excel_path.is_file():
            raise CommandError(f"No se encontró el archivo: {excel_path}")

        wb = load_workbook(excel_path, data_only=True)
        if "Documento_x_Criterio" not in wb.sheetnames:
            raise CommandError("La hoja 'Documento_x_Criterio' no existe en el Excel.")
        ws = wb["Documento_x_Criterio"]

        docs = {d.codigo: d for d in Documento.objects.all()}
        crits = {c.codigo: c for c in Criterio.objects.all()}

        created = 0
        existing = 0
        skipped = 0
        missing_docs: set[str] = set()
        missing_crits: set[str] = set()

        for row in ws.iter_rows(min_row=2, values_only=True):
            id_doc = (str(row[0]).strip() if row[0] is not None else "")
            id_crit = (str(row[2]).strip() if row[2] is not None else "")
            if not id_doc or not id_crit:
                continue

            documento = docs.get(id_doc)
            criterio = crits.get(id_crit)
            if documento is None:
                missing_docs.add(id_doc)
                skipped += 1
                continue
            if criterio is None:
                missing_crits.add(id_crit)
                skipped += 1
                continue

            _, was_created = DocumentoCriterio.objects.get_or_create(
                documento=documento,
                criterio=criterio,
            )
            if was_created:
                created += 1
            else:
                existing += 1

        self.stdout.write(
            self.style.SUCCESS(
                f"Importación lista: creados={created}, ya existían={existing}, omitidos={skipped}"
            )
        )
        if missing_docs:
            self.stdout.write(
                self.style.WARNING(
                    f"Documentos no encontrados ({len(missing_docs)}): "
                    f"{', '.join(sorted(missing_docs)[:15])}"
                    + ("..." if len(missing_docs) > 15 else "")
                )
            )
        if missing_crits:
            self.stdout.write(
                self.style.WARNING(
                    f"Criterios no encontrados ({len(missing_crits)}): "
                    f"{', '.join(sorted(missing_crits)[:15])}"
                    + ("..." if len(missing_crits) > 15 else "")
                )
            )
