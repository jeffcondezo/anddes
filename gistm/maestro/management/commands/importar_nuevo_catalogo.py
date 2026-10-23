from __future__ import annotations

import re
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from openpyxl import load_workbook

from maestro.models import (
    Actividad,
    Acuerdo,
    Area,
    AvisoConsultor,
    Criterio,
    Documento,
    DocumentoCarga,
    DocumentoCriterio,
    DocumentoMensaje,
    Empresa,
    EmpresaDocumento,
    EmpresaDocumentoCriterio,
    EmpresaDocumentoPregunta,
    EmpresaNotificacionRegla,
    PerfilUsuario,
    Pregunta,
    Principio,
    RegistroLoginExterno,
    Requisito,
    Revision,
    RevisionCriterioDesactivado,
    RolDocumentoCriterio,
    Tema,
)


def _principio_codigo(raw: str) -> str:
    match = re.search(r"(\d+)", raw or "")
    if not match:
        raise ValueError(f"No se pudo interpretar el principio: {raw!r}")
    return f"P{int(match.group(1)):02d}"


def _requisito_codigo(raw: str) -> str:
    match = re.search(r"(\d+)\.(\d+)", raw or "")
    if not match:
        raise ValueError(f"No se pudo interpretar el requisito: {raw!r}")
    return f"R{int(match.group(1)):02d}.{int(match.group(2))}"


def _extraer_preguntas(texto: str | None) -> list[str]:
    """Separa preguntas por ¿ … ? (incluye saltos de línea internos)."""
    if not texto:
        return []
    raw = str(texto).strip()
    if not raw:
        return []
    partes = re.findall(r"¿[^¿]*\?", raw, flags=re.S)
    limpias: list[str] = []
    for p in partes:
        t = re.sub(r"\s+", " ", p).strip()
        if t:
            limpias.append(t)
    return limpias


def _wipe_prueba(stdout, style) -> None:
    """Borra data operativa de prueba y catálogo GISTM (conserva usuarios)."""
    # Desvincular perfiles de empresas antes de borrar empresas.
    PerfilUsuario.objects.update(empresa=None)
    for perfil in PerfilUsuario.objects.all():
        perfil.empresas.clear()

    counts = {
        "EmpresaDocumentoPregunta": EmpresaDocumentoPregunta.objects.count(),
        "EmpresaDocumentoCriterio": EmpresaDocumentoCriterio.objects.count(),
        "RevisionCriterioDesactivado": RevisionCriterioDesactivado.objects.count(),
        "DocumentoMensaje": DocumentoMensaje.objects.count(),
        "DocumentoCarga": DocumentoCarga.objects.count(),
        "EmpresaDocumento": EmpresaDocumento.objects.count(),
        "AvisoConsultor": AvisoConsultor.objects.count(),
        "Acuerdo": Acuerdo.objects.count(),
        "Actividad": Actividad.objects.count(),
        "EmpresaNotificacionRegla": EmpresaNotificacionRegla.objects.count(),
        "RegistroLoginExterno": RegistroLoginExterno.objects.count(),
        "Area": Area.objects.count(),
        "Revision": Revision.objects.count(),
        "Empresa": Empresa.objects.count(),
        "DocumentoCriterio": DocumentoCriterio.objects.count(),
        "Pregunta": Pregunta.objects.count(),
        "Criterio": Criterio.objects.count(),
        "Requisito": Requisito.objects.count(),
        "Principio": Principio.objects.count(),
        "Tema": Tema.objects.count(),
        "Documento": Documento.objects.count(),
    }
    for name, n in counts.items():
        stdout.write(f"  wipe {name}: {n}")

    EmpresaDocumentoPregunta.objects.all().delete()
    EmpresaDocumentoCriterio.objects.all().delete()
    RevisionCriterioDesactivado.objects.all().delete()
    DocumentoMensaje.objects.all().delete()
    DocumentoCarga.objects.all().delete()
    EmpresaDocumento.objects.all().delete()
    AvisoConsultor.objects.all().delete()
    Acuerdo.objects.all().delete()
    Actividad.objects.all().delete()
    EmpresaNotificacionRegla.objects.all().delete()
    RegistroLoginExterno.objects.all().delete()
    Revision.objects.all().delete()
    Area.objects.all().delete()
    Empresa.objects.all().delete()

    DocumentoCriterio.objects.all().delete()
    Pregunta.objects.all().delete()
    Criterio.objects.all().delete()
    Requisito.objects.all().delete()
    Principio.objects.all().delete()
    Tema.objects.all().delete()
    Documento.objects.all().delete()

    stdout.write(style.WARNING("Data de prueba y catálogo anterior eliminados."))


class Command(BaseCommand):
    help = (
        "Borra data de prueba e importa catálogo desde nuevo.xlsx: "
        "temas/principios/requisitos/criterios, preguntas, documentos y relación estándar."
    )

    def add_arguments(self, parser):
        parser.add_argument(
            "--path",
            default="",
            help="Ruta al Excel (por defecto gistm/nuevo.xlsx).",
        )
        parser.add_argument(
            "--keep-data",
            action="store_true",
            help="No borrar data existente (solo upsert catálogo).",
        )

    @transaction.atomic
    def handle(self, *args, **options):
        path_opt = (options.get("path") or "").strip()
        if path_opt:
            excel_path = Path(path_opt)
        else:
            excel_path = Path(__file__).resolve().parents[3] / "nuevo.xlsx"

        if not excel_path.is_file():
            raise CommandError(f"No se encontró el archivo: {excel_path}")

        wb = load_workbook(excel_path, data_only=True)
        required = [
            "Criterios GISTM",
            "Criterios",
            "Catalogo_Documentos",
            "Relacion_Doc_Criterio",
        ]
        for sheet in required:
            if sheet not in wb.sheetnames:
                raise CommandError(f"Falta la hoja '{sheet}' en {excel_path.name}.")

        if not options.get("keep_data"):
            self.stdout.write("Limpiando data de prueba…")
            _wipe_prueba(self.stdout, self.style)

        # --- Temas / Principios / Requisitos / Criterios ---
        ws = wb["Criterios GISTM"]
        temas_cache: dict[str, Tema] = {}
        princ_cache: dict[str, Principio] = {}
        req_cache: dict[str, Requisito] = {}
        crit_created = 0
        crit_updated = 0

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            id_crit = str(row[0]).strip()
            tema_cod = str(row[1]).strip() if row[1] is not None else ""
            tema_nom = str(row[2]).strip() if row[2] is not None else ""
            princ_raw = str(row[3]).strip() if row[3] is not None else ""
            princ_desc = str(row[4]).strip() if row[4] is not None else ""
            req_raw = str(row[5]).strip() if row[5] is not None else ""
            req_desc = str(row[6]).strip() if row[6] is not None else ""
            crit_desc = str(row[8]).strip() if len(row) > 8 and row[8] is not None else ""
            if not id_crit or not tema_cod or not princ_raw or not req_raw:
                continue

            if tema_cod not in temas_cache:
                tema, _ = Tema.objects.update_or_create(
                    codigo=tema_cod,
                    defaults={"nombre": tema_nom or tema_cod, "activo": True},
                )
                temas_cache[tema_cod] = tema
            else:
                tema = temas_cache[tema_cod]
                if tema_nom and tema.nombre != tema_nom:
                    tema.nombre = tema_nom
                    tema.save(update_fields=["nombre", "actualizado_en"])

            p_cod = _principio_codigo(princ_raw)
            if p_cod not in princ_cache:
                princ, _ = Principio.objects.update_or_create(
                    codigo=p_cod,
                    defaults={
                        "descripcion": princ_desc or princ_raw,
                        "tema": tema,
                        "activo": True,
                    },
                )
                princ_cache[p_cod] = princ
            else:
                princ = princ_cache[p_cod]
                changed = False
                if princ_desc and princ.descripcion != princ_desc:
                    princ.descripcion = princ_desc
                    changed = True
                if princ.tema_id != tema.id:
                    princ.tema = tema
                    changed = True
                if changed:
                    princ.save()

            r_cod = _requisito_codigo(req_raw)
            if r_cod not in req_cache:
                req, _ = Requisito.objects.update_or_create(
                    codigo=r_cod,
                    defaults={
                        "descripcion": req_desc or req_raw,
                        "principio": princ,
                        "activo": True,
                    },
                )
                req_cache[r_cod] = req
            else:
                req = req_cache[r_cod]
                changed = False
                if req_desc and req.descripcion != req_desc:
                    req.descripcion = req_desc
                    changed = True
                if req.principio_id != princ.id:
                    req.principio = princ
                    changed = True
                if changed:
                    req.save()

            defaults = {
                "descripcion": crit_desc or id_crit,
                "requisito": req,
                "principio": princ,
                "tema": tema,
                "activo": True,
            }
            crit, created = Criterio.objects.update_or_create(
                codigo=id_crit,
                defaults=defaults,
            )
            if created:
                crit_created += 1
            else:
                crit_updated += 1

        # Recalcular ponderaciones 1/N por requisito
        for req in Requisito.objects.all():
            req.recalcular_ponderaciones_criterios()

        self.stdout.write(
            self.style.SUCCESS(
                f"Catálogo: {len(temas_cache)} temas, {len(princ_cache)} principios, "
                f"{len(req_cache)} requisitos, criterios +{crit_created}/~{crit_updated}."
            )
        )

        # --- Preguntas (hoja Criterios) ---
        ws = wb["Criterios"]
        criterios_by_code = {c.codigo: c for c in Criterio.objects.all()}
        codigo_corto_a_id: dict[str, str] = {}
        preg_created = 0
        preg_skipped_crit = 0
        preguntas_por_criterio: dict[str, int] = defaultdict(int)

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[1] is None:
                continue
            codigo_corto = str(row[0]).strip() if row[0] is not None else ""
            id_plat = str(row[1]).strip()
            if codigo_corto:
                codigo_corto_a_id[codigo_corto.lower()] = id_plat
            criterio = criterios_by_code.get(id_plat)
            if criterio is None:
                preg_skipped_crit += 1
                continue
            textos = _extraer_preguntas(row[7] if len(row) > 7 else None)
            for i, texto in enumerate(textos, start=1):
                preguntas_por_criterio[id_plat] += 1
                n = preguntas_por_criterio[id_plat]
                codigo_preg = f"{id_plat}-Q{n:02d}"
                _, created = Pregunta.objects.update_or_create(
                    codigo=codigo_preg,
                    defaults={
                        "texto": texto,
                        "criterio": criterio,
                        "orden": i,
                        "activo": True,
                    },
                )
                if created:
                    preg_created += 1

        self.stdout.write(
            self.style.SUCCESS(
                f"Preguntas: {Pregunta.objects.count()} total "
                f"(+{preg_created} nuevas; criterios sin match: {preg_skipped_crit})."
            )
        )

        # --- Documentos ---
        ws = wb["Catalogo_Documentos"]
        doc_created = 0
        doc_updated = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            codigo = str(row[0]).strip()
            nombre = str(row[1]).strip() if row[1] is not None else codigo
            tipo = str(row[2]).strip() if row[2] is not None else "Documento"
            if not codigo:
                continue
            _, created = Documento.objects.update_or_create(
                codigo=codigo,
                defaults={"nombre": nombre, "tipo_documento": tipo or "Documento"},
            )
            if created:
                doc_created += 1
            else:
                doc_updated += 1

        self.stdout.write(
            self.style.SUCCESS(
                f"Documentos: +{doc_created} / actualizados {doc_updated} "
                f"(total {Documento.objects.count()})."
            )
        )

        # --- Relación documento–criterio (principal + secundario + rol) ---
        ws = wb["Relacion_Doc_Criterio"]
        docs = {d.codigo: d for d in Documento.objects.all()}
        rel_created = 0
        rel_updated = 0
        skipped = 0
        missing_docs: set[str] = set()
        missing_crits: set[str] = set()
        con_secundario = 0

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            codigo_corto = str(row[0]).strip()
            rol_raw = str(row[2]).strip().upper() if len(row) > 2 and row[2] else ""
            id_doc = str(row[3]).strip() if len(row) > 3 and row[3] is not None else ""
            id_sec_raw = row[5] if len(row) > 5 else None
            # Algunas celdas traen varios códigos (salto de línea); usamos el primero válido.
            id_sec_candidates: list[str] = []
            if id_sec_raw is not None:
                for part in re.split(r"[\n,;]+", str(id_sec_raw)):
                    tok = part.strip()
                    if tok and tok not in ("-", "None", "none"):
                        id_sec_candidates.append(tok)
            id_crit = codigo_corto_a_id.get(codigo_corto.lower(), "")
            if not id_doc or not id_crit:
                skipped += 1
                continue
            documento = docs.get(id_doc)
            criterio = criterios_by_code.get(id_crit)
            if documento is None:
                missing_docs.add(id_doc)
                skipped += 1
                continue
            if criterio is None:
                missing_crits.add(id_crit)
                skipped += 1
                continue
            doc_sec = None
            for id_sec in id_sec_candidates:
                candidate = docs.get(id_sec)
                if candidate is None:
                    missing_docs.add(id_sec)
                    continue
                if candidate.pk == documento.pk:
                    continue
                doc_sec = candidate
                con_secundario += 1
                break

            if rol_raw.startswith("COMP"):
                rol = RolDocumentoCriterio.COMPLEMENTARIO
            else:
                rol = RolDocumentoCriterio.PRINCIPAL

            _, created = DocumentoCriterio.objects.update_or_create(
                criterio=criterio,
                defaults={
                    "documento": documento,
                    "documento_secundario": doc_sec,
                    "rol": rol,
                },
            )
            if created:
                rel_created += 1
            else:
                rel_updated += 1

        self.stdout.write(
            self.style.SUCCESS(
                f"Relación estándar: +{rel_created} / actualizadas {rel_updated} "
                f"(con secundario {con_secundario}, omitidas {skipped})."
            )
        )
        if missing_docs:
            self.stdout.write(
                self.style.WARNING(
                    f"Documentos faltantes en relación: {', '.join(sorted(missing_docs)[:20])}"
                )
            )
        if missing_crits:
            self.stdout.write(
                self.style.WARNING(
                    f"Criterios faltantes en relación: {', '.join(sorted(missing_crits)[:20])}"
                )
            )

        self.stdout.write(self.style.SUCCESS("Importación desde nuevo.xlsx completada."))
