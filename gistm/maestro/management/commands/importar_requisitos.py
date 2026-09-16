from __future__ import annotations

import re
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from maestro.models import Criterio, Principio, Requisito


def _principio_codigo(raw: str) -> str:
    """'Principio 01' / 'Principio 1' -> 'P01'."""
    match = re.search(r"(\d+)", raw or "")
    if not match:
        raise ValueError(f"No se pudo interpretar el principio: {raw!r}")
    return f"P{int(match.group(1)):02d}"


def _requisito_codigo(raw: str) -> str:
    """'Requisito 01.1' / 'Requisito 1.2' -> 'R01.1'."""
    match = re.search(r"(\d+)\.(\d+)", raw or "")
    if not match:
        raise ValueError(f"No se pudo interpretar el requisito: {raw!r}")
    return f"R{int(match.group(1)):02d}.{int(match.group(2))}"


class Command(BaseCommand):
    help = (
        "Importa requisitos desde la hoja 'Criterios GISTM' del Excel y "
        "reasigna cada criterio a su requisito."
    )

    def add_arguments(self, parser):
        parser.add_argument(
            "--path",
            default="",
            help="Ruta al excel.xlsx (por defecto gistm/excel.xlsx).",
        )

    @transaction.atomic
    def handle(self, *args, **options):
        path_opt = (options.get("path") or "").strip()
        if path_opt:
            excel_path = Path(path_opt)
        else:
            excel_path = Path(__file__).resolve().parents[3] / "excel.xlsx"

        if not excel_path.is_file():
            raise CommandError(f"No se encontró el archivo: {excel_path}")

        wb = load_workbook(excel_path, data_only=True)
        if "Criterios GISTM" not in wb.sheetnames:
            raise CommandError("La hoja 'Criterios GISTM' no existe en el Excel.")
        ws = wb["Criterios GISTM"]

        principios = {p.codigo: p for p in Principio.objects.all()}
        criterios = {
            c.codigo: c
            for c in Criterio.objects.select_related("requisito", "principio")
        }

        created_req = 0
        updated_req = 0
        linked = 0
        unchanged = 0
        missing_princ: set[str] = set()
        missing_crit: set[str] = set()
        skipped = 0
        seen_req_codes: set[str] = set()

        for row in ws.iter_rows(min_row=2, values_only=True):
            id_crit = (str(row[0]).strip() if row[0] is not None else "")
            princ_raw = (str(row[3]).strip() if row[3] is not None else "")
            req_raw = (str(row[5]).strip() if row[5] is not None else "")
            req_desc = (str(row[6]).strip() if row[6] is not None else "")

            if not id_crit or not req_raw:
                continue

            try:
                princ_codigo = _principio_codigo(princ_raw)
                req_codigo = _requisito_codigo(req_raw)
            except ValueError as exc:
                self.stdout.write(self.style.WARNING(str(exc)))
                skipped += 1
                continue

            principio = principios.get(princ_codigo)
            if principio is None:
                missing_princ.add(princ_codigo)
                skipped += 1
                continue

            requisito, was_created = Requisito.objects.update_or_create(
                codigo=req_codigo,
                defaults={
                    "descripcion": req_desc or req_raw,
                    "principio": principio,
                },
            )
            seen_req_codes.add(req_codigo)
            if was_created:
                created_req += 1
            else:
                updated_req += 1

            criterio = criterios.get(id_crit)
            if criterio is None:
                missing_crit.add(id_crit)
                skipped += 1
                continue

            if (
                criterio.requisito_id == requisito.id
                and criterio.principio_id == principio.id
            ):
                unchanged += 1
                continue

            criterio.requisito = requisito
            criterio.principio = principio
            criterio.tema = principio.tema
            criterio.save()
            linked += 1

        # Quitar placeholders de migración que ya no tienen criterios.
        orphans = Requisito.objects.filter(codigo__startswith="REQ-").exclude(
            criterios__isnull=False
        )
        deleted_placeholders, _ = orphans.delete()

        # Recalcular ponderaciones 1/N y sumas de principio.
        for requisito in Requisito.objects.select_related("principio").all():
            requisito.recalcular_ponderaciones_criterios()

        self.stdout.write(
            self.style.SUCCESS(
                "Importación de requisitos lista: "
                f"requisitos_creados={created_req}, "
                f"requisitos_actualizados={updated_req}, "
                f"criterios_reasignados={linked}, "
                f"criterios_sin_cambio={unchanged}, "
                f"omitidos={skipped}, "
                f"placeholders_eliminados={deleted_placeholders}, "
                f"requisitos_en_excel={len(seen_req_codes)}"
            )
        )
        if missing_princ:
            self.stdout.write(
                self.style.WARNING(
                    f"Principios no encontrados: {', '.join(sorted(missing_princ))}"
                )
            )
        if missing_crit:
            self.stdout.write(
                self.style.WARNING(
                    f"Criterios no encontrados ({len(missing_crit)}): "
                    f"{', '.join(sorted(missing_crit)[:20])}"
                    + ("..." if len(missing_crit) > 20 else "")
                )
            )
